import os
from datetime import date, datetime, timezone, time, timedelta
import json
import psycopg2
import sys
import time
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Border, Side, Alignment
from openpyxl.styles import PatternFill
from pytz import timezone
import calendar
from openpyxl.styles import PatternFill
import pandas as pd
from . import crud, schemas
from sqlalchemy.orm import Session
from enum import Enum
from openpyxl.styles import Font

import os
import sys

if getattr(sys, 'frozen', False):
    dirname = os.path.dirname(sys.executable)
else:
    dirname = os.path.dirname(os.path.abspath(__file__))


async def generate_report_shift_data(db: Session, date_: date):
    wb = load_workbook(os.path.join(dirname, f"SPHINX_REPORT_TEMPLATE.xlsx"))
    data_shift = await crud.get_report_data_by_date_(db, date_)

    ws = wb["Data"]

    try:
        for i, datap in enumerate(data_shift):
            print(datap)
            row = i + 3
            ws[f'A{row}'] = date_
            ws[f'B{row}'] = datap[10]
            ws[f'D{row}'] = datap[11]
            ws[f'E{row}'] = datap[0]
            ws[f'F{row}'] = datap[1]
            ws[f'G{row}'] = datap[2]
            ws[f'H{row}'] = datap[3]
            ws[f'I{row}'] = datap[4]
            ws[f'J{row}'] = datap[5]
            ws[f'K{row}'] = datap[6]
            ws[f'L{row}'] = datap[7]
            ws[f'M{row}'] = datap[8]
            ws[f'N{row}'] = datap[9]

            ws[f'A{row}'].alignment = Alignment(horizontal="center")
            ws[f'B{row}'].alignment = Alignment(horizontal="center")
            ws[f'C{row}'].alignment = Alignment(horizontal="center")
            ws[f'D{row}'].alignment = Alignment(horizontal="center")
            ws[f'E{row}'].alignment = Alignment(horizontal="center")
            ws[f'F{row}'].alignment = Alignment(horizontal="center")
            ws[f'G{row}'].alignment = Alignment(horizontal="center")
            ws[f'H{row}'].alignment = Alignment(horizontal="center")
            ws[f'I{row}'].alignment = Alignment(horizontal="center")
            ws[f'J{row}'].alignment = Alignment(horizontal="center")
            ws[f'K{row}'].alignment = Alignment(horizontal="center")
            ws[f'L{row}'].alignment = Alignment(horizontal="center")
            ws[f'M{row}'].alignment = Alignment(horizontal="center")
            ws[f'N{row}'].alignment = Alignment(horizontal="center")

            ws[f'A{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'B{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'C{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'D{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'E{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'F{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'G{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'H{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'I{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'J{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'K{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'L{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'M{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'N{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))

    except Exception as e:
        print(e)

    # path_ = os.path.join(dirname, f"Reports/ReportData_{date_}.xlsx")
    path_ = os.path.join(dirname, f"Reports/ReportData_{date_}.xlsx")
    wb.save(path_)
    return path_


# ...............................................................................................

async def generate_report_machine_data(db: Session, date_: date, machine_name: str):
    wb = load_workbook(os.path.join(dirname, f"SPHINX_REPORT_TEMPLATE.xlsx"))
    data_shift = await crud.get_report_machine_data(db, date_, machine_name)

    ws = wb["Data"]

    try:
        for i, datap in enumerate(data_shift):
            print(datap)
            row = i + 3
            ws[f'A{row}'] = date_
            ws[f'B{row}'] = machine_name
            ws[f'D{row}'] = datap[10]
            ws[f'E{row}'] = datap[0]
            ws[f'F{row}'] = datap[1]
            ws[f'G{row}'] = datap[2]
            ws[f'H{row}'] = datap[3]
            ws[f'I{row}'] = datap[4]
            ws[f'J{row}'] = datap[5]
            ws[f'K{row}'] = datap[6]
            ws[f'L{row}'] = datap[7]
            ws[f'M{row}'] = datap[8]
            ws[f'N{row}'] = datap[9]

            ws[f'A{row}'].alignment = Alignment(horizontal="center")
            ws[f'B{row}'].alignment = Alignment(horizontal="center")
            ws[f'C{row}'].alignment = Alignment(horizontal="center")
            ws[f'D{row}'].alignment = Alignment(horizontal="center")
            ws[f'E{row}'].alignment = Alignment(horizontal="center")
            ws[f'F{row}'].alignment = Alignment(horizontal="center")
            ws[f'G{row}'].alignment = Alignment(horizontal="center")
            ws[f'H{row}'].alignment = Alignment(horizontal="center")
            ws[f'I{row}'].alignment = Alignment(horizontal="center")
            ws[f'J{row}'].alignment = Alignment(horizontal="center")
            ws[f'K{row}'].alignment = Alignment(horizontal="center")
            ws[f'L{row}'].alignment = Alignment(horizontal="center")
            ws[f'M{row}'].alignment = Alignment(horizontal="center")
            ws[f'N{row}'].alignment = Alignment(horizontal="center")

            ws[f'A{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'B{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'C{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'D{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'E{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'F{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'G{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'H{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'I{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'J{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'K{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'L{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'M{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))
            ws[f'N{row}'].border = Border(left=Side(border_style="thin"),
                                          right=Side(border_style="thin"),
                                          top=Side(border_style="thin"),
                                          bottom=Side(border_style="thin"))

    except Exception as e:
        print(e)

    # path_ = os.path.join(dirname, f"Reports/ReportData_{date_}.xlsx")
    path_ = os.path.join(dirname, f"Reports/ReportData_{date_}.xlsx")
    wb.save(path_)
    return path_


def generate_daily_report(date_: date):
    date_from = (date_).strftime("%F")

    conn = psycopg2.connect(database="thingsboard", user="postgres", password="Cybershot#903", host='localhost',
                            port=5432)
    print(conn)
    cur = conn.cursor()
    cur.execute(f'''SELECT id,name,type from device WHERE type in ('HAAS','Fanuc', 'MAKINO', 'Siemens');''', "\n")
    devices_sansera = cur.fetchall()
    print("devices_sansera", devices_sansera)
    wb = load_workbook(os.path.join(dirname, f"SPHINX_REPORT_TEMPLATE.xlsx"))

    for idx, (id_, device, type) in enumerate(devices_sansera):
        device = device.replace("/", "-")
        device = device.replace("*", "-")
        print(device)
        ws = wb["Data"]
        row = (idx + 3) + (2 * idx)
        ws[f'A{row}'] = date_from
        ws[f'B{row}'] = device
        ws[f'C{row}'] = type
        ws[f'D{row}'] = "A"
        print(device, len(device))

        a = (idx + 3) + (2 * idx)
        b = (idx + 5) + (2 * idx)
        ws.merge_cells(f'A{a}:A{b}')
        ws.merge_cells(f'B{a}:B{b}')
        ws.merge_cells(f'C{a}:C{b}')

        row2 = (idx + 4) + (2 * idx)
        ws[f'D{row2}'] = "B"

        row3 = (idx + 5) + (2 * idx)
        ws[f'D{row3}'] = "C"

        row_c = (6 * idx) + 3
        row_d = (6 * idx) + 6

        start_color = "00FF00"  # green
        end_color = "DDFFDD"  # light green

        # fill = PatternFill(start_color='FFE6E6E6', end_color="FFBFA9", fill_type='solid')
        # alt_fill = PatternFill(start_color='FFFFFFFF', end_color="DDDDDD", fill_type='solid')

        # ws[f'A{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")
        # ws[f'B{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")
        # ws[f'C{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")

        if row % 2 == 0:
            ws[f'A{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")
            ws[f'B{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")
            ws[f'C{row}'].fill = PatternFill(start_color=start_color, end_color="FFBFA9", fill_type="gray125")
        else:
            ws[f'A{row}'].fill = PatternFill(start_color=start_color, end_color="DDDDDD", fill_type="gray125")
            ws[f'B{row}'].fill = PatternFill(start_color=start_color, end_color="DDDDDD", fill_type="gray125")
            ws[f'C{row}'].fill = PatternFill(start_color=start_color, end_color="DDDDDD", fill_type="gray125")

        ws[f'D{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'D{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'D{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'E{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'E{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'E{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'F{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'F{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'F{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'G{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'G{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'G{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'H{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'H{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'H{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'I{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'I{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'I{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'J{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'J{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'J{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'K{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'K{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'K{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'L{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'L{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'L{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'M{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'M{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'M{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")
        ws[f'N{row}'].fill = PatternFill(start_color=start_color, end_color="C9EEFF", fill_type="gray125")
        ws[f'N{row2}'].fill = PatternFill(start_color=start_color, end_color="97DEFF", fill_type="gray125")
        ws[f'N{row3}'].fill = PatternFill(start_color=start_color, end_color="62CDFF", fill_type="gray125")

        ws[f'A{row}'].alignment = Alignment(horizontal="center", vertical='center')
        ws[f'B{row}'].alignment = Alignment(horizontal="center", vertical='center')
        ws[f'C{row}'].alignment = Alignment(horizontal="center", vertical='center')
        ws[f'D{row}'].alignment = Alignment(horizontal="center")
        ws[f'D{row2}'].alignment = Alignment(horizontal="center")
        ws[f'D{row3}'].alignment = Alignment(horizontal="center")
        ws[f'E{row}'].alignment = Alignment(horizontal="center")
        ws[f'E{row2}'].alignment = Alignment(horizontal="center")
        ws[f'E{row3}'].alignment = Alignment(horizontal="center")
        ws[f'F{row}'].alignment = Alignment(horizontal="center")
        ws[f'F{row2}'].alignment = Alignment(horizontal="center")
        ws[f'F{row3}'].alignment = Alignment(horizontal="center")
        ws[f'G{row}'].alignment = Alignment(horizontal="center")
        ws[f'G{row2}'].alignment = Alignment(horizontal="center")
        ws[f'G{row3}'].alignment = Alignment(horizontal="center")
        ws[f'H{row}'].alignment = Alignment(horizontal="center")
        ws[f'H{row2}'].alignment = Alignment(horizontal="center")
        ws[f'H{row3}'].alignment = Alignment(horizontal="center")
        ws[f'I{row}'].alignment = Alignment(horizontal="center")
        ws[f'I{row2}'].alignment = Alignment(horizontal="center")
        ws[f'I{row3}'].alignment = Alignment(horizontal="center")
        ws[f'J{row}'].alignment = Alignment(horizontal="center")
        ws[f'J{row2}'].alignment = Alignment(horizontal="center")
        ws[f'J{row3}'].alignment = Alignment(horizontal="center")
        ws[f'K{row}'].alignment = Alignment(horizontal="center")
        ws[f'K{row2}'].alignment = Alignment(horizontal="center")
        ws[f'K{row3}'].alignment = Alignment(horizontal="center")
        ws[f'L{row}'].alignment = Alignment(horizontal="center")
        ws[f'L{row2}'].alignment = Alignment(horizontal="center")
        ws[f'L{row3}'].alignment = Alignment(horizontal="center")
        ws[f'M{row}'].alignment = Alignment(horizontal="center")
        ws[f'M{row2}'].alignment = Alignment(horizontal="center")
        ws[f'M{row3}'].alignment = Alignment(horizontal="center")
        ws[f'N{row}'].alignment = Alignment(horizontal="center")
        ws[f'N{row2}'].alignment = Alignment(horizontal="center")
        ws[f'N{row3}'].alignment = Alignment(horizontal="center")

        ws[f'A{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'A{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"))

        ws[f'A{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'B{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'B{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"))

        ws[f'B{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'C{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'C{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"))

        ws[f'C{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'D{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'D{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'D{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'E{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'E{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'E{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'F{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'F{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'F{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'G{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'G{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'G{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'H{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'H{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'H{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'I{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'I{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'I{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'J{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'J{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'J{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'K{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'K{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'K{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'L{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'L{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'L{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'M{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'M{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'M{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        ws[f'N{row}'].border = Border(left=Side(border_style="thin"),
                                      right=Side(border_style="thin"),
                                      top=Side(border_style="thin"),
                                      bottom=Side(border_style="thin"))
        ws[f'N{row2}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))
        ws[f'N{row3}'].border = Border(left=Side(border_style="thin"),
                                       right=Side(border_style="thin"),
                                       top=Side(border_style="thin"),
                                       bottom=Side(border_style="thin"))

        # ....................................SHIFT - A ...............................................

        start_ts_a = date_from + " 06:30:00"  # SHIFT -A
        stop_ts_a = date_from + " 14:59:59"
        local_tz = timezone('Asia/Kolkata')  ## Set your timezone
        start_ts_a = datetime.strptime(start_ts_a, "%Y-%m-%d %H:%M:%S")
        stop_ts_a = datetime.strptime(stop_ts_a, "%Y-%m-%d %H:%M:%S")
        start_ts_a = local_tz.normalize(local_tz.localize(start_ts_a, is_dst=False)).timestamp() * 1000
        stop_ts_a = local_tz.normalize(local_tz.localize(stop_ts_a, is_dst=False)).timestamp() * 1000
        # print("unix", start_ts_a, stop_ts_a)
        print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_a) / 1000))
        print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_a) / 1000))

        # ....................................SHIFT - B ...............................................

        start_ts_b = date_from + " 15:00:00"  # SHIFT - B
        stop_ts_b = date_from + " 23:29:59"
        local_tz = timezone('Asia/Kolkata')  ## Set your timezone
        start_ts_b = datetime.strptime(start_ts_b, "%Y-%m-%d %H:%M:%S")
        stop_ts_b = datetime.strptime(stop_ts_b, "%Y-%m-%d %H:%M:%S")
        start_ts_b = local_tz.normalize(local_tz.localize(start_ts_b, is_dst=False)).timestamp() * 1000
        stop_ts_b = local_tz.normalize(local_tz.localize(stop_ts_b, is_dst=False)).timestamp() * 1000
        # print("unix", start_ts_b, stop_ts_b)
        print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_b) / 1000))
        print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_b) / 1000))

        # ....................................SHIFT - C ...............................................

        start_ts_c = date_from + " 23:30:00"  # SHIFT -C
        stop_ts_c = date_from + " 06:29:59"
        local_tz = timezone('Asia/Kolkata')  ## Set your timezone
        start_ts_c = datetime.strptime(start_ts_c, "%Y-%m-%d %H:%M:%S")
        stop_ts_c = datetime.strptime(stop_ts_c, "%Y-%m-%d %H:%M:%S") + timedelta(days=1)
        start_ts_c = local_tz.normalize(local_tz.localize(start_ts_c, is_dst=False)).timestamp() * 1000
        stop_ts_c = local_tz.normalize(local_tz.localize(stop_ts_c, is_dst=False)).timestamp() * 1000
        # print("unix", start_ts_c, stop_ts_c)
        print("START_DATE & TIME Shift_C:-->", datetime.fromtimestamp((start_ts_c) / 1000))
        print("STOP DATE & TIME shift_C :-->", datetime.fromtimestamp((stop_ts_c) / 1000))

        # ............................SHIFT-A.....................................................

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                              AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Target_Count = cur.fetchall()
            print("Target_Count : ", Target_Count)
            Target_Count1 = [x[0] for x in Target_Count]
            print(Target_Count1)
            row = (idx + 3) + (2 * idx)
            if Target_Count1[-1] < 0:
                ws[f'E{row}'] = 0
            else:
                ws[f'E{row}'] = Target_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                              AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Part_Count = cur.fetchall()
            print("Part_Count : ", Part_Count)
            Part_Count1 = [x[0] for x in Part_Count]
            print(Part_Count1)
            row = (idx + 3) + (2 * idx)
            if Part_Count1[-1] < 0:
                ws[f'F{row}'] = 0
            else:
                ws[f'F{row}'] = Part_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                                  AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            rejection = cur.fetchall()
            print("	rejection : ", rejection)
            rejection1 = [x[0] for x in rejection]
            print(rejection1)
            row = (idx + 3) + (2 * idx)
            if rejection1 == [] or rejection1[-1] < 0:
                ws[f'G{row}'] = 0
            else:
                ws[f'G{row}'] = rejection1[-1]

        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                          AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            OEE = cur.fetchall()
            print("OEE : ", OEE)
            OEE1 = [x[0] for x in OEE]
            print(OEE1)
            print("last", OEE1[-1])
            row = (idx + 3) + (2 * idx)
            if OEE1[-1] < 0:
                ws[f'H{row}'] = 0
            else:
                ws[f'H{row}'] = OEE1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                          AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Efficiency = cur.fetchall()
            print("Efficiency : ", Efficiency)
            Efficiency1 = [x[0] for x in Efficiency]
            print(Efficiency1)
            print("last", Efficiency1[-1])
            row = (idx + 3) + (2 * idx)
            if Efficiency1[-1] == None or Efficiency1[-1] == [] or Efficiency1[-1] < 0:
                ws[f'I{row}'] = 0
            else:
                ws[f'I{row}'] = Efficiency1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                          AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Quality = cur.fetchall()
            print("Quality : ", Quality)
            Quality1 = [x[0] for x in Quality]
            print(Quality1)
            row = (idx + 3) + (2 * idx)
            if Quality1[-1] < 0:
                ws[f'J{row}'] = 0
            else:
                ws[f'J{row}'] = Quality1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                           AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                           AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Availability = cur.fetchall()
            print("Availability : ", Availability)
            Availability1 = [x[0] for x in Availability]
            print(Availability1)
            row = (idx + 3) + (2 * idx)
            if Availability1[-1] < 0:
                ws[f'K{row}'] = 0
            else:
                ws[f'K{row}'] = Availability1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                          AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Idle_Time = cur.fetchall()
            print("	Idle_Time : ", Idle_Time)
            Idle_Time1 = [x[0] for x in Idle_Time]
            print(Idle_Time1)
            row = (idx + 3) + (2 * idx)
            if Idle_Time1[-1] < 0:
                ws[f'L{row}'] = 0
            else:
                ws[f'L{row}'] = Idle_Time1[-1]

        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                      AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                      AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Breakdown_Time = cur.fetchall()
            print("Breakdown_Time : ", Breakdown_Time)
            Breakdown_Time1 = [x[0] for x in Breakdown_Time]
            print(Breakdown_Time1)
            row = (idx + 3) + (2 * idx)
            if Breakdown_Time1[-1] < 0:
                ws[f'M{row}'] = 0
            else:
                ws[f'M{row}'] = Breakdown_Time1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                          AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
            Breakdown_Count = cur.fetchall()
            print("Breakdown_Count : ", Breakdown_Count)
            Breakdown_Count1 = [x[0] for x in Breakdown_Count]
            print(Breakdown_Count1)
            row = (idx + 3) + (2 * idx)
            if Breakdown_Count1[-1] < 0:
                ws[f'N{row}'] = 0
            else:
                ws[f'N{row}'] = Breakdown_Count1[-1]
        except Exception as e:
            print(e)

        # .................................SHIFT - B..........................................................

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Target_Count = cur.fetchall()
            print("Target_Count : ", Target_Count)
            Target_Count1 = [x[0] for x in Target_Count]
            print(Target_Count1)
            row = (idx + 4) + (2 * idx)
            if Target_Count1[-1] < 0:
                ws[f'E{row}'] = 0
            else:
                ws[f'E{row}'] = Target_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Part_Count = cur.fetchall()
            print("Part_Count : ", Part_Count)
            Part_Count1 = [x[0] for x in Part_Count]
            print(Part_Count1)
            row = (idx + 4) + (2 * idx)
            if Part_Count1[-1] < 0:
                ws[f'F{row}'] = 0
            else:
                ws[f'F{row}'] = Part_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                                  AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            rejection = cur.fetchall()
            print("	rejection : ", rejection)
            rejection1 = [x[0] for x in rejection]
            print(rejection1)
            row = (idx + 4) + (2 * idx)
            if rejection1 == [] or rejection1[-1] < 0:
                ws[f'G{row}'] = 0
            else:
                ws[f'G{row}'] = rejection1[-1]

        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            OEE = cur.fetchall()
            print("OEE : ", OEE)
            OEE1 = [x[0] for x in OEE]
            print(OEE1)
            print("last", OEE1[-1])
            row = (idx + 4) + (2 * idx)
            if OEE1[-1] < 0:
                ws[f'H{row}'] = 0
            else:
                ws[f'H{row}'] = OEE1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Efficiency = cur.fetchall()
            print("Efficiency : ", Efficiency)
            Efficiency1 = [x[0] for x in Efficiency]
            print(Efficiency1)
            print("last", Efficiency1[-1])
            row = (idx + 4) + (2 * idx)
            if Efficiency1[-1] == None or Efficiency1[-1] == [] or Efficiency1[-1] < 0:
                ws[f'I{row}'] = 0
            else:
                ws[f'I{row}'] = Efficiency1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Quality = cur.fetchall()
            print("Quality : ", Quality)
            Quality1 = [x[0] for x in Quality]
            print(Quality1)
            row = (idx + 4) + (2 * idx)
            if Quality1[-1] < 0:
                ws[f'J{row}'] = 0
            else:
                ws[f'J{row}'] = Quality1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                           AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                           AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Availability = cur.fetchall()
            print("Availability : ", Availability)
            Availability1 = [x[0] for x in Availability]
            print(Availability1)
            row = (idx + 4) + (2 * idx)
            if Availability1[-1] < 0:
                ws[f'K{row}'] = 0
            else:
                ws[f'K{row}'] = Availability1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Idle_Time = cur.fetchall()
            print("	Idle_Time : ", Idle_Time)
            Idle_Time1 = [x[0] for x in Idle_Time]
            print(Idle_Time1)
            row = (idx + 4) + (2 * idx)
            if Idle_Time1[-1] < 0:
                ws[f'L{row}'] = 0
            else:
                ws[f'L{row}'] = Idle_Time1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                      AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                      AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Breakdown_Time = cur.fetchall()
            print("Breakdown_Time : ", Breakdown_Time)
            Breakdown_Time1 = [x[0] for x in Breakdown_Time]
            print(Breakdown_Time1)
            row = (idx + 4) + (2 * idx)
            if Breakdown_Time1[-1] < 0:
                ws[f'M{row}'] = 0
            else:
                ws[f'M{row}'] = Breakdown_Time1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
            Breakdown_Count = cur.fetchall()
            print("Breakdown_Count : ", Breakdown_Count)
            Breakdown_Count1 = [x[0] for x in Breakdown_Count]
            print(Breakdown_Count1)
            row = (idx + 4) + (2 * idx)
            if Breakdown_Count1[-1] < 0:
                ws[f'N{row}'] = 0
            else:
                ws[f'N{row}'] = Breakdown_Count1[-1]
        except Exception as e:
            print(e)

        # ......................................SHIFT - C....................................................

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Target_Count = cur.fetchall()
            print("Target_Count : ", Target_Count)
            Target_Count1 = [x[0] for x in Target_Count]
            print(Target_Count1)
            row = (idx + 5) + (2 * idx)
            if Target_Count1[-1] < 0:
                ws[f'E{row}'] = 0
            else:
                ws[f'E{row}'] = Target_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                            AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                             AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Part_Count = cur.fetchall()
            print("Part_Count : ", Part_Count)
            Part_Count1 = [x[0] for x in Part_Count]
            print(Part_Count1)

            row = (idx + 5) + (2 * idx)
            if Part_Count1[-1] < 0:
                ws[f'F{row}'] = 0
            else:
                ws[f'F{row}'] = Part_Count1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                             AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                             AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            rejection = cur.fetchall()
            print("	rejection : ", rejection)
            rejection1 = [x[0] for x in rejection]
            print(rejection1)
            row = (idx + 5) + (2 * idx)
            if rejection1 == [] or rejection1[-1] < 0:
                ws[f'G{row}'] = 0
            else:
                ws[f'G{row}'] = rejection1[-1]

        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            OEE = cur.fetchall()
            print("OEE : ", OEE)
            OEE1 = [x[0] for x in OEE]
            print(OEE1)
            print("last", OEE1[-1])
            row = (idx + 5) + (2 * idx)
            if OEE1[-1] < 0:
                ws[f'H{row}'] = 0
            else:
                ws[f'H{row}'] = OEE1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                          AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Efficiency = cur.fetchall()
            print("Efficiency : ", Efficiency)
            Efficiency1 = [x[0] for x in Efficiency]
            print(Efficiency1)
            print("last", Efficiency1[-1])
            row = (idx + 5) + (2 * idx)
            if Efficiency1[-1] == None or Efficiency1[-1] == [] or Efficiency1[-1] < 0:
                ws[f'I{row}'] = 0
            else:
                ws[f'I{row}'] = Efficiency1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Quality = cur.fetchall()
            print("Quality : ", Quality)
            Quality1 = [x[0] for x in Quality]
            print(Quality1)
            row = (idx + 5) + (2 * idx)
            if Quality1[-1] < 0:
                ws[f'J{row}'] = 0
            else:
                ws[f'J{row}'] = Quality1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                                  AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Availability = cur.fetchall()
            print("Availability : ", Availability)
            Availability1 = [x[0] for x in Availability]
            print(Availability1)
            row = (idx + 5) + (2 * idx)
            if Availability1[-1] < 0:
                ws[f'K{row}'] = 0
            else:
                ws[f'K{row}'] = Availability1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Idle_Time = cur.fetchall()
            print("	Idle_Time : ", Idle_Time)
            Idle_Time1 = [x[0] for x in Idle_Time]
            print(Idle_Time1)
            row = (idx + 5) + (2 * idx)
            if Idle_Time1[-1] < 0:
                ws[f'L{row}'] = 0
            else:
                ws[f'L{row}'] = Idle_Time1[-1]


        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                             AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                             AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Breakdown_Time = cur.fetchall()
            print("Breakdown_Time : ", Breakdown_Time)
            Breakdown_Time1 = [x[0] for x in Breakdown_Time]
            print(Breakdown_Time1)
            row = (idx + 5) + (2 * idx)
            if Breakdown_Time1[-1] < 0:
                ws[f'M{row}'] = 0
            else:
                ws[f'M{row}'] = Breakdown_Time1[-1]
        except Exception as e:
            print(e)

        try:
            cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
            Breakdown_Count = cur.fetchall()
            print("Breakdown_Count : ", Breakdown_Count)
            Breakdown_Count1 = [x[0] for x in Breakdown_Count]
            print(Breakdown_Count1)
            row = (idx + 5) + (2 * idx)
            if Breakdown_Count1[-1] < 0:
                ws[f'N{row}'] = 0
            else:
                ws[f'N{row}'] = Breakdown_Count1[-1]
        except Exception as e:
            print(e)

    conn.close()

    path_ = os.path.join(dirname, f"Reports/Sphinx_Enterprise_Reports_{date_}.xlsx")
    wb.save(path_)
    return path_


# generate_daily_report(datetime(2023, 3, 1))

# ...............................................................................................................


def generate_daily_data(date_: date, shift: str, machine_name: str):
    date_from = (date_).strftime("%F")
    print("shift", shift)

    conn = psycopg2.connect(database="thingsboard", user="postgres", password="Cybershot#903", host='localhost',
                            port=5432)
    print(conn)
    cur = conn.cursor()

    cur.execute(f'''SELECT id,name from device WHERE name = '{machine_name}';''', "\n")
    devices_sansera = cur.fetchall()
    print(devices_sansera)
    result = {}
    result1 = {}

    for idx, (id_, device) in enumerate(devices_sansera):
        print("machine name :", device)
        device = device.replace("/", "-")
        device = device.replace("*", "-")
        print(device)
        if shift == "A":
            start_ts_a = date_from + " 06:30:00"  # SHIFT -A
            stop_ts_a = date_from + " 14:59:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_a = datetime.strptime(start_ts_a, "%Y-%m-%d %H:%M:%S")
            stop_ts_a = datetime.strptime(stop_ts_a, "%Y-%m-%d %H:%M:%S")
            start_ts_a = local_tz.normalize(local_tz.localize(start_ts_a, is_dst=False)).timestamp() * 1000
            stop_ts_a = local_tz.normalize(local_tz.localize(stop_ts_a, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_a, stop_ts_a)
            print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_a) / 1000))
            print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_a) / 1000))

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                                  AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Target_Count = cur.fetchall()
                print("Target_Count : ", Target_Count)
                Target_Count1 = [x[0] for x in Target_Count]
                print(Target_Count1)
                TC = Target_Count1[-1]
                # print("Target_Count : ", TC)

                result["Target_Count"] = TC

            except Exception as e:
                result["Target_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                                  AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Part_Count = cur.fetchall()
                # print("Part_Count : ", Part_Count)
                Part_Count1 = [x[0] for x in Part_Count]
                # print(Part_Count1)
                PC = Part_Count1[-1]
                # print("Part_Count : ", PC)
                result["Part_Count"] = PC
            except Exception as e:
                result["Part_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                      AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                                      AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                rejection = cur.fetchall()
                # print("	rejection : ", rejection)
                rejection1 = [x[0] for x in rejection]
                # print(rejection1)
                rej = rejection1[-1]
                # print("rejection : ", rej)
                result["Rejection"] = rej


            except Exception as e:
                result["Rejection"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                              AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                OEE = cur.fetchall()
                # print("OEE : ", OEE)
                OEE1 = [x[0] for x in OEE]
                # print(OEE1)
                # print("last", OEE1[-1])
                oee = OEE1[-1]
                # print("OEE :", oee)
                result["OEE"] = oee

            except Exception as e:
                result["OEE"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                              AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Efficiency = cur.fetchall()
                # print("Efficiency : ", Efficiency)
                Efficiency1 = [x[0] for x in Efficiency]
                # print(Efficiency1)
                # print("last", Efficiency1[-
                eff = Efficiency1[-1]
                # print("Performance :", eff)
                result["Performance"] = eff

            except Exception as e:
                result["Performance"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                              AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Quality = cur.fetchall()
                # print("Quality : ", Quality)
                Quality1 = [x[0] for x in Quality]
                # print(Quality1)
                quality = Quality1[-1]
                # print("Quality :", quality)
                result["Quality"] = quality

            except Exception as e:
                result["Quality"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                               AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                               AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Availability = cur.fetchall()
                # print("Availability : ", Availability)
                Availability1 = [x[0] for x in Availability]
                # print(Availability1)
                avail = Availability1[-1]
                # print("Availability :", avail)
                result["Availability"] = avail

            except Exception as e:
                result["Availability"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                        AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                                        AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Idle_Time = cur.fetchall()
                # print("	Idle_Time : ", Idle_Time)
                Idle_Time1 = [x[0] for x in Idle_Time]
                # print(Idle_Time1)
                idl = Idle_Time1[-1]
                # print("Idle_Time :", idl)
                result["Idle_Time"] = idl
            except Exception as e:
                result["Idle_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                    AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                                    AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Breakdown_Time = cur.fetchall()
                # print("Breakdown_Time : ", Breakdown_Time)
                Breakdown_Time1 = [x[0] for x in Breakdown_Time]
                # print(Breakdown_Time1)
                Br_t = Breakdown_Time1[-1]
                print("Breakdown_time : ", Br_t)
                result["Breakdown_Time"] = Br_t

            except Exception as e:
                result["Breakdown_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                        AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                                        AND ts > {start_ts_a} AND ts < {stop_ts_a};''')
                Breakdown_Count = cur.fetchall()
                # print("Breakdown_Count : ", Breakdown_Count)
                Breakdown_Count1 = [x[0] for x in Breakdown_Count]
                # print(Breakdown_Count1)
                Br_c = Breakdown_Count1[-1]
                # print("Breakdown_Count : ", Br_c)
                result["Breakdown_Count"] = Br_c

            except Exception as e:
                result["Breakdown_Count"] = ""
                print(e)

            for key, val in result.items():
                if val == None:
                    y = ""
                    result.update({key: y})
            # print("result", result)
            conn.close()
            return result

        # ....................................SHIFT - B ...............................................
        elif shift == "B":

            start_ts_b = date_from + " 15:00:00"  # SHIFT - B
            stop_ts_b = date_from + " 23:29:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_b = datetime.strptime(start_ts_b, "%Y-%m-%d %H:%M:%S")
            stop_ts_b = datetime.strptime(stop_ts_b, "%Y-%m-%d %H:%M:%S")
            start_ts_b = local_tz.normalize(local_tz.localize(start_ts_b, is_dst=False)).timestamp() * 1000
            stop_ts_b = local_tz.normalize(local_tz.localize(stop_ts_b, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_b, stop_ts_b)
            print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_b) / 1000))
            print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_b) / 1000))

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                                  AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Target_Count = cur.fetchall()
                # print("Target_Count : ", Target_Count)
                Target_Count1 = [x[0] for x in Target_Count]
                # print(Target_Count1)
                TC = Target_Count1[-1]
                # print("Target_Count : ", TC)
                result["Target_Count"] = TC
            except Exception as e:
                result["Target_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                  AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                                  AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Part_Count = cur.fetchall()
                # print("Part_Count : ", Part_Count)
                Part_Count1 = [x[0] for x in Part_Count]
                # print(Part_Count1)
                PC = Part_Count1[-1]
                # print("Part_Count : ", PC)
                result["Part_Count"] = PC
            except Exception as e:
                result["Part_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                      AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                                      AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                rejection = cur.fetchall()
                # print("	rejection : ", rejection)
                rejection1 = [x[0] for x in rejection]
                # print(rejection1)
                rej = rejection1[-1]
                # print("rejection : ", rej)
                result["Rejection"] = rej


            except Exception as e:
                result["Rejection"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                OEE = cur.fetchall()
                # print("OEE : ", OEE)
                OEE1 = [x[0] for x in OEE]
                # print(OEE1)
                # print("last", OEE1[-1])
                oee = OEE1[-1]
                # print("OEE :", oee)
                result["OEE"] = oee

            except Exception as e:
                result["OEE"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Efficiency = cur.fetchall()
                # print("Efficiency : ", Efficiency)
                Efficiency1 = [x[0] for x in Efficiency]
                # print(Efficiency1)
                # print("last", Efficiency1[-
                eff = Efficiency1[-1]
                # print("Performance :", eff)
                result["Performance"] = eff
            except Exception as e:
                result["Performance"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Quality = cur.fetchall()
                # print("Quality : ", Quality)
                Quality1 = [x[0] for x in Quality]
                # print(Quality1)
                quality = Quality1[-1]
                # print("Quality :", quality)
                result["Quality"] = quality
            except Exception as e:
                result["Quality"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                               AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                               AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Availability = cur.fetchall()
                # print("Availability : ", Availability)
                Availability1 = [x[0] for x in Availability]
                # print(Availability1)
                avail = Availability1[-1]
                # print("Availability :", avail)
                result["Availability"] = avail

            except Exception as e:
                result["Availability"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Idle_Time = cur.fetchall()
                # print("	Idle_Time : ", Idle_Time)
                Idle_Time1 = [x[0] for x in Idle_Time]
                # print(Idle_Time1)
                idl = Idle_Time1[-1]
                # print("Idle_Time :", idl)
                result["Idle_Time"] = idl
            except Exception as e:
                result["Idle_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                          AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                          AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Breakdown_Time = cur.fetchall()
                # print("Breakdown_Time : ", Breakdown_Time)
                Breakdown_Time1 = [x[0] for x in Breakdown_Time]
                # print(Breakdown_Time1)
                Br_t = Breakdown_Time1[-1]
                # print("Breaktime_time : ", Br_t)
                result["Breakdown_Time"] = Br_t


            except Exception as e:
                result["Breakdown_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                              AND ts > {start_ts_b} AND ts < {stop_ts_b};''')
                Breakdown_Count = cur.fetchall()
                # print("Breakdown_Count : ", Breakdown_Count)
                Breakdown_Count1 = [x[0] for x in Breakdown_Count]
                # print(Breakdown_Count1)
                Br_c = Breakdown_Count1[-1]
                # print("Breakdown_Count : ", Br_c)
                result["Breakdown_Count"] = Br_c


            except Exception as e:
                result["Breakdown_Count"] = ""
                print(e)

            for key, val in result.items():
                if val == None:
                    y = ""
                    result.update({key: y})
            print("result", result)
            conn.close()
            return result

        # ....................................SHIFT - C ...............................................
        else:

            start_ts_c = date_from + " 23:30:00"  # SHIFT -C
            stop_ts_c = date_from + " 06:29:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_c = datetime.strptime(start_ts_c, "%Y-%m-%d %H:%M:%S")
            stop_ts_c = datetime.strptime(stop_ts_c, "%Y-%m-%d %H:%M:%S") + timedelta(days=1)
            start_ts_c = local_tz.normalize(local_tz.localize(start_ts_c, is_dst=False)).timestamp() * 1000
            stop_ts_c = local_tz.normalize(local_tz.localize(stop_ts_c, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_c, stop_ts_c)
            print("START_DATE & TIME Shift_C:-->", datetime.fromtimestamp((start_ts_c) / 1000))
            print("STOP DATE & TIME shift_C :-->", datetime.fromtimestamp((stop_ts_c) / 1000))

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                     AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Target_Count' LIMIT 1)
                                     AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Target_Count = cur.fetchall()
                # print("Target_Count : ", Target_Count)
                Target_Count1 = [x[0] for x in Target_Count]
                # print(Target_Count1)
                TC = Target_Count1[-1]
                # print("Target_Count : ", TC)
                result["Target_Count"] = TC
            except Exception as e:
                result["Target_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Part_Count' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Part_Count = cur.fetchall()
                # print("Part_Count : ", Part_Count)
                Part_Count1 = [x[0] for x in Part_Count]
                # print(Part_Count1)
                PC = Part_Count1[-1]
                # print("Part_Count : ", PC)
                result["Part_Count"] = PC

            except Exception as e:
                result["Part_Count"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='rejection' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                rejection = cur.fetchall()
                # print("	rejection : ", rejection)
                rejection1 = [x[0] for x in rejection]
                # print(rejection1)
                rej = rejection1[-1]
                # print("rejection : ", rej)
                result["Rejection"] = rej

            except Exception as e:
                result["Rejection"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                     AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='OEE' LIMIT 1)
                                     AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                OEE = cur.fetchall()
                # print("OEE : ", OEE)
                OEE1 = [x[0] for x in OEE]
                # print(OEE1)
                # print("last", OEE1[-1])
                oee = OEE1[-1]
                # print("OEE :", oee)
                result["OEE"] = oee


            except Exception as e:
                result["OEE"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                              AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Efficiency' LIMIT 1)
                              AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Efficiency = cur.fetchall()
                # print("Efficiency : ", Efficiency)
                Efficiency1 = [x[0] for x in Efficiency]
                # print(Efficiency1)
                # print("last", Efficiency1[-
                eff = Efficiency1[-1]
                # print("Performance :", eff)
                result["Performance"] = eff
            except Exception as e:
                result["Performance"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                     AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Quality' LIMIT 1)
                                     AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Quality = cur.fetchall()
                # print("Quality : ", Quality)
                Quality1 = [x[0] for x in Quality]
                # print(Quality1)
                quality = Quality1[-1]
                # print("Quality :", quality)
                result["Quality"] = quality

            except Exception as e:
                result["Quality"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                      AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Availability' LIMIT 1)
                                      AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Availability = cur.fetchall()
                # print("Availability : ", Availability)
                Availability1 = [x[0] for x in Availability]
                # print(Availability1)
                avail = Availability1[-1]
                # print("Availability :", avail)
                result["Availability"] = avail

            except Exception as e:
                result["Availability"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                     AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Idle_Time' LIMIT 1)
                                     AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Idle_Time = cur.fetchall()
                # print("	Idle_Time : ", Idle_Time)
                Idle_Time1 = [x[0] for x in Idle_Time]
                # print(Idle_Time1)
                idl = Idle_Time1[-1]
                # print("Idle_Time :", idl)
                result["Idle_Time"] = idl


            except Exception as e:
                result["Idle_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT dbl_v from ts_kv WHERE entity_Id='{id_}'
                                 AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Time' LIMIT 1)
                                 AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Breakdown_Time = cur.fetchall()
                # print("Breakdown_Time : ", Breakdown_Time)
                Breakdown_Time1 = [x[0] for x in Breakdown_Time]
                # print(Breakdown_Time1)
                Br_t = Breakdown_Time1[-1]
                # print("Breaktime_time : ", Br_t)
                result["Breakdown_Time"] = Br_t

            except Exception as e:
                result["Breakdown_Time"] = ""
                print(e)

            try:
                cur.execute(f'''SELECT long_v from ts_kv WHERE entity_Id='{id_}'
                                     AND key=(SELECT key_id FROM ts_kv_dictionary WHERE key='Breakdown_Count' LIMIT 1)
                                     AND ts > {start_ts_c} AND ts < {stop_ts_c};''')
                Breakdown_Count = cur.fetchall()
                # print("Breakdown_Count : ", Breakdown_Count)
                Breakdown_Count1 = [x[0] for x in Breakdown_Count]
                # print(Breakdown_Count1)
                Br_c = Breakdown_Count1[-1]
                print("Breakdown_Count : ", Br_c)
                result["Breakdown_Count"] = Br_c

            except Exception as e:
                result["Breakdown_Count"] = ""
                print(e)
            # print(result)

            for key, val in result.items():
                if val == None:
                    y = ""
                    result.update({key: y})
            # print("result", result)
            conn.close()
            return result



# generate_daily_data(datetime(2023, 3, 1), "A", 'MACR-0015-VMC-436')


# ................................Downtime Report........................................................

def generate_downtime_report(date_: date, shift: str, machine_name: str):
    date_to = (date_ + timedelta(days=1)).strftime("%F")
    date_from = date_.strftime("%F")
    conn = psycopg2.connect(database="thingsboard", user="postgres", password="Cybershot#903", host='localhost',
                            port=5432)
    print(conn)
    cur = conn.cursor()
    # cur.execute(f'''SELECT id,name from device WHERE name = '{machine_name}';''', "\n")
    cur.execute(f'''SELECT id,name from device WHERE name = '{machine_name}';''', "\n")
    devices = cur.fetchall()
    print(devices, len(devices))
    wb = load_workbook(os.path.join(dirname, f"Sphinx_Downtime_Report.xlsx"))

    # sheet_template = wb['Temp']
    ws = wb['Temp']
    for idx, (id_, device) in enumerate(devices):
        # if device not in machine_names:
        #     continue
        device = device.replace("/", "-")
        device = device.replace("*", "-")
        ws.title = device
        if shift == 'A':

            # ....................................SHIFT - A ...............................................

            start_ts_a = date_from + " 06:30:00"  # SHIFT -A
            stop_ts_a = date_from + " 14:59:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_a = datetime.strptime(start_ts_a, "%Y-%m-%d %H:%M:%S")
            stop_ts_a = datetime.strptime(stop_ts_a, "%Y-%m-%d %H:%M:%S")
            start_ts_a = local_tz.normalize(local_tz.localize(start_ts_a, is_dst=False)).timestamp() * 1000
            stop_ts_a = local_tz.normalize(local_tz.localize(stop_ts_a, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_a, stop_ts_a)
            print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_a) / 1000))
            print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_a) / 1000))

            try:
                cur.execute(f'''SELECT start_ts, end_ts, additional_info, (end_ts-start_ts)/60000
                                from alarm WHERE originator_id='{id_}' AND start_ts > {start_ts_a} AND
                                end_ts < {stop_ts_a} AND  (end_ts-start_ts) > 0 ORDER BY start_ts;''')
                downtime_data = cur.fetchall()
                print("downtime_data", downtime_data)
                downtime_data2 = []

                for i, (start_ts, end_ts, additional_info, duration) in enumerate(downtime_data):
                    if additional_info == "null" or additional_info == "{}":
                        type_ = ""
                        reason = ""
                    else:
                        print(type(json.loads(additional_info)), json.loads(additional_info))
                        type_ = json.loads(additional_info)['reason']
                        reason = json.loads(additional_info)['category']
                    dur_list = [start_ts, end_ts, duration, type_, reason]
                    downtime_data2.append(dur_list)

                downtime_df = pd.DataFrame(downtime_data2,
                                           columns=["start_ts", "end_ts", "duration", "type_", "reason"])
                downtime_df['start_ts'] = pd.to_datetime(downtime_df['start_ts'], utc=True, unit='ms')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_localize(None)
                downtime_df['end_ts'] = pd.to_datetime(downtime_df['end_ts'], utc=True, unit='ms')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_localize(None)
                print(downtime_df.head())

                for i, (start_ts, end_ts, duration, type_, reason) in enumerate(downtime_df.values.tolist()[:]):
                    row = i + 3
                    ws[f'A{row}'] = start_ts.strftime("%Y-%m-%d")
                    ws[f'B{row}'] = start_ts.strftime("%H:%M:%S")
                    ws[f'C{row}'] = end_ts.strftime("%H:%M:%S")
                    ws[f'D{row}'] = duration
                    ws[f'E{row}'] = type_
                    ws[f'F{row}'] = reason

                    ws[f'A{row}'].alignment = Alignment(horizontal="center")
                    ws[f'B{row}'].alignment = Alignment(horizontal="center")
                    ws[f'C{row}'].alignment = Alignment(horizontal="center")
                    ws[f'D{row}'].alignment = Alignment(horizontal="center")
                    ws[f'E{row}'].alignment = Alignment(horizontal="center")
                    ws[f'F{row}'].alignment = Alignment(horizontal="center")

                data_df = downtime_df.iloc[:, [2, 4]]
                print(data_df)
                df = data_df.groupby('reason', as_index=False).sum()
                print(df)
                for i, datap in enumerate(df.values.tolist()[:]):
                    row = i + 3
                    if datap[0] == "":
                        ws[f'H{row}'] = "Undefined"
                        ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    else:
                        ws[f'H{row}'] = datap[0]
                    ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'] = datap[1]

            except Exception as e:
                print(e)

        # ....................................SHIFT - B ...............................................
        elif shift == "B":

            start_ts_b = date_from + " 15:00:00"  # SHIFT - B
            stop_ts_b = date_from + " 23:29:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_b = datetime.strptime(start_ts_b, "%Y-%m-%d %H:%M:%S")
            stop_ts_b = datetime.strptime(stop_ts_b, "%Y-%m-%d %H:%M:%S")
            start_ts_b = local_tz.normalize(local_tz.localize(start_ts_b, is_dst=False)).timestamp() * 1000
            stop_ts_b = local_tz.normalize(local_tz.localize(stop_ts_b, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_b, stop_ts_b)
            print("START_DATE & TIME:-->", datetime.fromtimestamp((start_ts_b) / 1000))
            print("STOP DATE & TIME:-->", datetime.fromtimestamp((stop_ts_b) / 1000))

            try:
                cur.execute(f'''SELECT start_ts, end_ts, additional_info, (end_ts-start_ts)/60000
                                from alarm WHERE originator_id='{id_}' AND start_ts > {start_ts_b} AND
                                end_ts < {stop_ts_b} AND  (end_ts-start_ts) > 0 ORDER BY start_ts;''')
                downtime_data = cur.fetchall()
                print("downtime_data", downtime_data)
                downtime_data2 = []

                for i, (start_ts, end_ts, additional_info, duration) in enumerate(downtime_data):
                    if additional_info == "null" or additional_info == "{}":
                        type_ = ""
                        reason = ""
                    else:
                        print(type(json.loads(additional_info)), json.loads(additional_info))
                        type_ = json.loads(additional_info)['reason']
                        reason = json.loads(additional_info)['category']
                    dur_list = [start_ts, end_ts, duration, type_, reason]
                    downtime_data2.append(dur_list)

                downtime_df = pd.DataFrame(downtime_data2,
                                           columns=["start_ts", "end_ts", "duration", "type_", "reason"])
                downtime_df['start_ts'] = pd.to_datetime(downtime_df['start_ts'], utc=True, unit='ms')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_localize(None)
                downtime_df['end_ts'] = pd.to_datetime(downtime_df['end_ts'], utc=True, unit='ms')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_localize(None)
                print(downtime_df.head())

                for i, (start_ts, end_ts, duration, type_, reason) in enumerate(downtime_df.values.tolist()[:]):
                    row = i + 3
                    ws[f'A{row}'] = start_ts.strftime("%Y-%m-%d")
                    ws[f'B{row}'] = start_ts.strftime("%H:%M:%S")
                    ws[f'C{row}'] = end_ts.strftime("%H:%M:%S")
                    ws[f'D{row}'] = duration
                    ws[f'E{row}'] = type_
                    ws[f'F{row}'] = reason

                    ws[f'A{row}'].alignment = Alignment(horizontal="center")
                    ws[f'B{row}'].alignment = Alignment(horizontal="center")
                    ws[f'C{row}'].alignment = Alignment(horizontal="center")
                    ws[f'D{row}'].alignment = Alignment(horizontal="center")
                    ws[f'E{row}'].alignment = Alignment(horizontal="center")
                    ws[f'F{row}'].alignment = Alignment(horizontal="center")

                data_df = downtime_df.iloc[:, [2, 4]]
                print(data_df)
                df = data_df.groupby('reason', as_index=False).sum()
                print(df)
                for i, datap in enumerate(df.values.tolist()[:]):
                    row = i + 3
                    if datap[0] == "":
                        ws[f'H{row}'] = "Undefined"
                        ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    else:
                        ws[f'H{row}'] = datap[0]
                    ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'] = datap[1]

            except Exception as e:
                print(e)

        else:

            start_ts_c = date_from + " 23:30:00"  # SHIFT -C
            stop_ts_c = date_from + " 06:29:59"
            local_tz = timezone('Asia/Kolkata')  ## Set your timezone
            start_ts_c = datetime.strptime(start_ts_c, "%Y-%m-%d %H:%M:%S")
            stop_ts_c = datetime.strptime(stop_ts_c, "%Y-%m-%d %H:%M:%S") + timedelta(days=1)
            start_ts_c = local_tz.normalize(local_tz.localize(start_ts_c, is_dst=False)).timestamp() * 1000
            stop_ts_c = local_tz.normalize(local_tz.localize(stop_ts_c, is_dst=False)).timestamp() * 1000
            # print("unix", start_ts_c, stop_ts_c)
            print("START_DATE & TIME Shift_C:-->", datetime.fromtimestamp((start_ts_c) / 1000))
            print("STOP DATE & TIME shift_C :-->", datetime.fromtimestamp((stop_ts_c) / 1000))

            try:
                cur.execute(f'''SELECT start_ts, end_ts, additional_info, (end_ts-start_ts)/60000
                                from alarm WHERE originator_id='{id_}' AND start_ts > {start_ts_c} AND
                                end_ts < {stop_ts_c} AND  (end_ts-start_ts) > 0 ORDER BY start_ts;''')
                downtime_data = cur.fetchall()
                print("downtime_data", downtime_data)
                downtime_data2 = []

                for i, (start_ts, end_ts, additional_info, duration) in enumerate(downtime_data):
                    if additional_info == "null" or additional_info == "{}":
                        type_ = ""
                        reason = ""
                    else:
                        print(type(json.loads(additional_info)), json.loads(additional_info))
                        type_ = json.loads(additional_info)['reason']
                        reason = json.loads(additional_info)['category']
                    dur_list = [start_ts, end_ts, duration, type_, reason]
                    downtime_data2.append(dur_list)

                downtime_df = pd.DataFrame(downtime_data2,
                                           columns=["start_ts", "end_ts", "duration", "type_", "reason"])
                downtime_df['start_ts'] = pd.to_datetime(downtime_df['start_ts'], utc=True, unit='ms')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['start_ts'] = downtime_df['start_ts'].dt.tz_localize(None)
                downtime_df['end_ts'] = pd.to_datetime(downtime_df['end_ts'], utc=True, unit='ms')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_convert('Asia/Kolkata')
                downtime_df['end_ts'] = downtime_df['end_ts'].dt.tz_localize(None)
                print(downtime_df.head())

                for i, (start_ts, end_ts, duration, type_, reason) in enumerate(downtime_df.values.tolist()[:]):
                    row = i + 3
                    ws[f'A{row}'] = start_ts.strftime("%Y-%m-%d")
                    ws[f'B{row}'] = start_ts.strftime("%H:%M:%S")
                    ws[f'C{row}'] = end_ts.strftime("%H:%M:%S")
                    ws[f'D{row}'] = duration
                    ws[f'E{row}'] = type_
                    ws[f'F{row}'] = reason

                    ws[f'A{row}'].alignment = Alignment(horizontal="center")
                    ws[f'B{row}'].alignment = Alignment(horizontal="center")
                    ws[f'C{row}'].alignment = Alignment(horizontal="center")
                    ws[f'D{row}'].alignment = Alignment(horizontal="center")
                    ws[f'E{row}'].alignment = Alignment(horizontal="center")
                    ws[f'F{row}'].alignment = Alignment(horizontal="center")

                data_df = downtime_df.iloc[:, [2, 4]]
                print(data_df)
                df = data_df.groupby('reason', as_index=False).sum()
                print(df)
                for i, datap in enumerate(df.values.tolist()[:]):
                    row = i + 3
                    if datap[0] == "":
                        ws[f'H{row}'] = "Undefined"
                        ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    else:
                        ws[f'H{row}'] = datap[0]
                    ws[f'H{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'].alignment = Alignment(horizontal="center")
                    ws[f'I{row}'] = datap[1]

            except Exception as e:
                print(e)

    conn.close()
    # wb.save("Daily.xlsx")

    # path_ = os.path.join(dirname, f"Reports/DowntimeReport_{date_}.xlsx")
    path_ = os.path.join(dirname, f"Reports/DowntimeReport_{date_}.xlsx")
    wb.save(path_)
    return path_


# .................................................Weighted Shift Report.........................................


async def generate_shift_report(db: Session, from_date: date, end_date: date, machine_name: str,
                                Shift: schemas.ShiftEnum):
    wb = load_workbook(os.path.join(dirname, f"SPHINX_HISTORICAL_REPORT_TEMPLATE.xlsx"))
    data_shift = await crud.get_report_shift_data(db, from_date, end_date, Shift, machine_name)
    data = await crud.get_combine_weighted_mean(db, from_date, end_date, Shift, machine_name)
    print(data)

    ws = wb["Data"]

    try:
        for i, datap in enumerate(data_shift):
            print(datap)
            row = i + 3

            ws[f'A{row}'] = datap['date_']
            ws[f'B{row}'] = machine_name
            ws[f'C{row}'] = datap['Shift']
            ws[f'D{row}'] = datap['Target_Count']
            ws[f'E{row}'] = datap['Actual_Target_Count']
            ws[f'F{row}'] = datap['Part_Count']
            ws[f'G{row}'] = datap['Rejection_Count']
            ws[f'H{row}'] = datap['Total_Time']
            ws[f'I{row}'] = datap['Operating_Time']
            ws[f'J{row}'] = datap['Breakdown_Time']
            ws[f'K{row}'] = datap['Idle_Time']
            ws[f'L{row}'] = datap['Breakdown_Count']
            ws[f'M{row}'] = datap['Efficiency']
            ws[f'N{row}'] = datap['Availability']
            ws[f'O{row}'] = datap['Quality']
            ws[f'P{row}'] = datap['OEE']

            columns = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P']
            thin_border = Border(left=Side(border_style="thin"),
                                 right=Side(border_style="thin"),
                                 top=Side(border_style="thin"),
                                 bottom=Side(border_style="thin"))

            for col in columns:
                cell = f'{col}{row}'
                ws[cell].alignment = Alignment(horizontal="center")
                ws[cell].border = thin_border

            row += 2

        ws[f'A{row}'] = 'Summary'

        ws[f'D{row}'] = data['Target_Count']
        ws[f'E{row}'] = data['Actual_Target_Count']
        ws[f'F{row}'] = data['Part_Count']
        ws[f'G{row}'] = data['Rejection_Count']
        ws[f'H{row}'] = data['Total_Time']
        ws[f'I{row}'] = data['Operating_Time']
        ws[f'J{row}'] = data['Breakdown_Time']
        ws[f'K{row}'] = data['Idle_Time']
        ws[f'L{row}'] = data['Breakdown_Count']
        ws[f'M{row}'] = data['Efficiency']
        ws[f'N{row}'] = data['Availability']
        ws[f'O{row}'] = data['Quality']
        ws[f'P{row}'] = data['OEE']

        columns = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P']

        center_alignment = Alignment(horizontal="center")
        thin_border = Border(left=Side(border_style="thin"),
                             right=Side(border_style="thin"),
                             top=Side(border_style="thin"),
                             bottom=Side(border_style="thin"))

        for col in columns:
            cell = f'{col}{row}'
            ws[cell].alignment = center_alignment
            ws[cell].border = thin_border
    except Exception as e:
        print(e)

    # path_ = os.path.join(dirname, f"Reports/ReportData_{from_date}.xlsx")
    path_ = os.path.join(dirname, f"Reports/ReportData_{from_date}.xlsx")
    wb.save(path_)
    return path_

# generate_downtime_report(datetime(2023, 3, 4), "C", 'SFB-402-0i-MD')
